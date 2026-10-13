from pathlib import Path
from shutil import copy2

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor as PptRGB
from pptx.util import Inches as PptInches, Pt as PptPt


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs" / "interview"
PNG_OUT = DOCS / "LeadFlow_ERD.png"
DOCX_SOURCE = DOCS / "LeadFlow_Technical_Case_Study.docx"
DOCX_OUT = DOCS / "LeadFlow_Technical_Case_Study_with_ERD.docx"
PPTX_SOURCE = DOCS / "LeadFlow_Technical_Presentation.pptx"
PPTX_OUT = DOCS / "LeadFlow_Technical_Presentation_with_ERD.pptx"


TABLES = {
    "users": [
        ("PK", "id", "BIGINT"), ("UQ", "email", "VARCHAR(254)"),
        ("", "name", "VARCHAR(120)"), ("", "password_hash", "VARCHAR(255)"),
        ("", "role", "ENUM(admin, sales)"), ("", "is_active", "BOOLEAN"),
        ("", "token_version", "INT"), ("", "created_at / updated_at", "TIMESTAMP"),
        ("", "deleted_at", "TIMESTAMP NULL"),
    ],
    "leads": [
        ("PK", "id", "BIGINT"), ("UQ", "phone", "VARCHAR(20) NULL"),
        ("UQ", "email", "VARCHAR(254) NULL"), ("", "name", "VARCHAR(150)"),
        ("", "source", "VARCHAR(80)"), ("", "status", "ENUM(6 states)"),
        ("FK", "assigned_user_id", "BIGINT NULL"), ("", "version", "INT"),
        ("FK", "created_by", "BIGINT"), ("FK", "updated_by", "BIGINT"),
        ("", "created_at / updated_at", "TIMESTAMP"), ("", "deleted_at", "TIMESTAMP NULL"),
    ],
    "enquiries": [
        ("PK", "id", "BIGINT"), ("FK", "lead_id", "BIGINT"),
        ("FK", "created_by", "BIGINT NULL"), ("", "source", "VARCHAR(80)"),
        ("", "raw_payload", "JSON NULL"), ("", "received_at", "TIMESTAMP"),
    ],
    "assignments": [
        ("PK", "id", "BIGINT"), ("FK", "lead_id", "BIGINT"),
        ("FK", "from_user_id", "BIGINT NULL"), ("FK", "to_user_id", "BIGINT"),
        ("FK", "assigned_by", "BIGINT"), ("", "reason", "VARCHAR(500) NULL"),
        ("", "assigned_at", "TIMESTAMP"), ("", "unassigned_at", "TIMESTAMP NULL"),
    ],
    "activities": [
        ("PK", "id", "BIGINT"), ("FK", "lead_id", "BIGINT"),
        ("FK", "actor_user_id", "BIGINT NULL"), ("", "type", "ENUM(8 types)"),
        ("", "message", "TEXT NULL"), ("", "changes", "JSON NULL"),
        ("", "created_at", "TIMESTAMP"),
    ],
    "reassignment_requests": [
        ("PK", "id", "BIGINT"), ("FK", "lead_id", "BIGINT"),
        ("FK", "requested_by", "BIGINT"), ("FK", "requested_to", "BIGINT NULL"),
        ("FK", "reviewed_by", "BIGINT NULL"), ("", "reason", "VARCHAR(500)"),
        ("", "status", "ENUM(4 states)"), ("", "reviewed_at", "TIMESTAMP NULL"),
        ("UQ", "pending_requester_id", "GENERATED"), ("", "created_at / updated_at", "TIMESTAMP"),
    ],
    "idempotency_requests": [
        ("PK", "id", "BIGINT"), ("FK", "user_id", "BIGINT"),
        ("UQ", "operation + key", "VARCHAR"), ("", "request_hash", "CHAR(64)"),
        ("", "status", "ENUM(2 states)"), ("", "response_status", "SMALLINT NULL"),
        ("", "response_body", "JSON NULL"), ("", "created / completed / expires", "TIMESTAMP"),
    ],
    "notifications": [
        ("PK", "id", "BIGINT"), ("FK", "user_id", "BIGINT"),
        ("FK", "lead_id", "BIGINT NULL"), ("", "type", "VARCHAR(60)"),
        ("", "title", "VARCHAR(160)"), ("", "message", "VARCHAR(500)"),
        ("", "is_read", "BOOLEAN"), ("", "created_at / read_at", "TIMESTAMP"),
    ],
}


def font(size, bold=False):
    candidates = [
        Path("C:/Windows/Fonts/seguisb.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf"),
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def rounded_line(draw, points, fill, width=5):
    draw.line(points, fill=fill, width=width, joint="curve")
    radius = width // 2
    for x, y in points[1:-1]:
        draw.ellipse((x-radius, y-radius, x+radius, y+radius), fill=fill)


def make_png():
    width, height = 3200, 1800
    image = Image.new("RGB", (width, height), "#F7FAFC")
    draw = ImageDraw.Draw(image)
    navy, blue, cyan = "#102A43", "#1769AA", "#00A6A6"
    ink, muted, border = "#243B53", "#627D98", "#BCCCDC"
    draw.rectangle((0, 0, width, 142), fill=navy)
    draw.text((90, 35), "LeadFlow CRM — Entity Relationship Diagram", font=font(54, True), fill="white")
    draw.text((90, 101), "MySQL 8.4  •  Current schema from server/v1/db/migrations", font=font(24), fill="#D9EAF7")

    positions = {
        "users": (80, 235, 660, 745),
        "idempotency_requests": (80, 1085, 660, 1580),
        "leads": (890, 355, 1570, 1050),
        "enquiries": (1830, 175, 2390, 580),
        "assignments": (2520, 175, 3120, 670),
        "activities": (1830, 755, 2390, 1205),
        "reassignment_requests": (2520, 790, 3120, 1395),
        "notifications": (1510, 1290, 2150, 1725),
    }

    # Relationships are routed behind entity cards. Each line means parent 1 → child many.
    rel = "#6B7C93"
    def route(points):
        rounded_line(draw, points, rel, 5)
        x, y = points[-1]
        draw.ellipse((x-8, y-8, x+8, y+8), fill=rel)

    route([(660, 365), (785, 365), (785, 470), (890, 470)])  # users -> leads
    route([(660, 525), (760, 525), (760, 1270), (370, 1270), (370, 1085)])  # users -> idem
    route([(660, 610), (810, 610), (810, 1505), (1510, 1505)])  # users -> notifications
    route([(1570, 490), (1700, 490), (1700, 375), (1830, 375)])  # leads -> enquiries
    route([(1570, 610), (2450, 610), (2450, 415), (2520, 415)])  # leads -> assignments
    route([(1570, 760), (1705, 760), (1705, 960), (1830, 960)])  # leads -> activities
    route([(1570, 885), (2405, 885), (2405, 1090), (2520, 1090)])  # leads -> reassign
    route([(1570, 990), (1630, 990), (1630, 1415), (1830, 1415)])  # leads -> notifications
    # Secondary user references into child tables.
    route([(660, 300), (2430, 300), (2430, 300), (2520, 300)])
    route([(660, 440), (1740, 440), (1740, 280), (1830, 280)])
    route([(660, 680), (1760, 680), (1760, 855), (1830, 855)])
    route([(660, 705), (2460, 705), (2460, 925), (2520, 925)])

    title_font, row_font, type_font, key_font = font(29, True), font(21), font(19), font(17, True)
    for name, rows in TABLES.items():
        x1, y1, x2, y2 = positions[name]
        core = name in {"users", "leads"}
        header = blue if core else cyan
        draw.rounded_rectangle((x1, y1, x2, y2), radius=22, fill="white", outline=border, width=4)
        draw.rounded_rectangle((x1, y1, x2, y1+68), radius=22, fill=header)
        draw.rectangle((x1, y1+46, x2, y1+68), fill=header)
        draw.text((x1+24, y1+16), name.upper(), font=title_font, fill="white")
        available = y2 - (y1 + 82)
        row_h = min(48, max(37, available // len(rows)))
        y = y1 + 84
        for i, (marker, field, dtype) in enumerate(rows):
            if i % 2:
                draw.rectangle((x1+3, y-3, x2-3, y+row_h-3), fill="#F3F7FA")
            if marker:
                color = {"PK": "#D64545", "FK": "#1769AA", "UQ": "#8A4FFF"}[marker]
                draw.rounded_rectangle((x1+18, y+3, x1+61, y+29), radius=6, fill=color)
                draw.text((x1+24, y+5), marker, font=key_font, fill="white")
            draw.text((x1+72, y+3), field, font=row_font, fill=ink)
            bbox = draw.textbbox((0, 0), dtype, font=type_font)
            draw.text((x2-18-(bbox[2]-bbox[0]), y+5), dtype, font=type_font, fill=muted)
            y += row_h

    # Relationship/cardinality legend.
    draw.rounded_rectangle((75, 1625, 1335, 1745), radius=18, fill="#E9F2F9", outline="#9FB3C8", width=3)
    draw.text((100, 1645), "Relationship notation", font=font(24, True), fill=navy)
    draw.text((100, 1688), "Parent  1", font=font(21), fill=ink)
    draw.line((225, 1703, 425, 1703), fill=rel, width=4)
    draw.text((440, 1688), "many  Child     •     optional foreign keys are marked NULL", font=font(21), fill=ink)
    draw.text((2230, 1665), "PK  primary key     FK  foreign key     UQ  unique", font=font(21, True), fill=ink)
    image.save(PNG_OUT, quality=95, dpi=(220, 220))


def add_docx():
    copy2(DOCX_SOURCE, DOCX_OUT)
    doc = Document(DOCX_OUT)
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11.69), Inches(8.27)
    section.top_margin = section.bottom_margin = Inches(0.45)
    section.left_margin = section.right_margin = Inches(0.55)

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("Appendix A — Database Entity Relationship Diagram")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(16, 42, 67)
    note = doc.add_paragraph("Current MySQL schema: 8 tables. Lines show one-to-many relationships from parent records to child records; nullable foreign keys represent optional relationships.")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_after = Pt(6)
    note.runs[0].font.size = Pt(9)
    note.runs[0].font.color.rgb = RGBColor(98, 125, 152)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(PNG_OUT), width=Inches(10.45))
    source = doc.add_paragraph("Source of truth: server/v1/db/migrations/001–009. Migration 008 adds the unique pending-request constraint and does not create a separate table.")
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source.runs[0].italic = True
    source.runs[0].font.size = Pt(8)
    source.runs[0].font.color.rgb = RGBColor(98, 125, 152)
    doc.save(DOCX_OUT)


def add_pptx():
    copy2(PPTX_SOURCE, PPTX_OUT)
    prs = Presentation(PPTX_OUT)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid(); bg.fore_color.rgb = PptRGB(247, 250, 252)

    banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, PptInches(0.78))
    banner.fill.solid(); banner.fill.fore_color.rgb = PptRGB(16, 42, 67)
    banner.line.fill.background()
    box = slide.shapes.add_textbox(PptInches(0.55), PptInches(0.12), PptInches(12.2), PptInches(0.5))
    tf = box.text_frame; tf.clear()
    p = tf.paragraphs[0]; p.text = "Database design"; p.font.name = "Aptos Display"; p.font.size = PptPt(26); p.font.bold = True; p.font.color.rgb = PptRGB(255,255,255)
    slide.shapes.add_picture(str(PNG_OUT), PptInches(1.15), PptInches(0.94), height=PptInches(6.2))

    # Put the ERD immediately after the auditability slide (existing slide 10).
    slide_ids = prs.slides._sldIdLst
    new_id = slide_ids[-1]
    slide_ids.remove(new_id)
    slide_ids.insert(10, new_id)
    prs.save(PPTX_OUT)


if __name__ == "__main__":
    DOCS.mkdir(parents=True, exist_ok=True)
    make_png()
    add_docx()
    add_pptx()
    print(PNG_OUT)
    print(DOCX_OUT)
    print(PPTX_OUT)
